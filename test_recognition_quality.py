"""Synthetic regression suite; no external AI, personal data or production writes."""
import datetime as dt
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import doc_reader
import training_catalog as cat
from scheduling import schedule,parse_dates
from production_calendar import is_working_day,span

class RecognitionQuality(unittest.TestCase):
    def test_calendar_totals(self):
        for year,total in [(2025,247),(2026,247),(2027,247)]:
            start=dt.date(year,1,1)
            self.assertEqual(sum(is_working_day(start+dt.timedelta(days=i)) for i in range(365)),total)
    def test_transferred_saturday(self):
        self.assertTrue(is_working_day(dt.date(2027,2,20)))
        self.assertFalse(is_working_day(dt.date(2027,2,22)))
        self.assertEqual(span(dt.date(2027,2,19),3),dt.date(2027,2,24))
    def test_focus_group_sequence(self):
        p=cat.resolve_many('высоте 2 группа и ОЗП 2 группа')
        r=schedule(p,'Дата приема 24.08.2026')
        self.assertEqual([x[1] for x in r],['24.08.2026 - 26.08.2026','27.08.2026 - 31.08.2026'])
    def test_hours_options(self):
        p=cat.resolve('ПП')
        self.assertEqual(schedule([p],'24.08.2026','start')[0][0]['days'],2)
        self.assertEqual(schedule([p],'24.08.2026','start',{p['id']:8})[0][0]['days'],1)
        self.assertEqual(schedule([p],'24.08.2026','start',{p['id']:178})[0][1],'')
    def test_worker_rank_duration(self):
        p=cat.resolve('Монтажник технологических трубопроводов 5 разряда')
        self.assertTrue(p['is_canonical'])
        self.assertEqual(p['hours'],180)
        self.assertEqual(schedule([p],'24.08.2026','start')[0][0]['days'],23)
    def test_missing_group_requires_review(self):
        for name in ('Высота', 'ОЗП', 'Пожарка', 'ПК 178 часов'):
            self.assertFalse(cat.resolve(name)['is_canonical'])
    def test_header_course_for_managers(self):
        p=doc_reader.extract_letter_programs(['Просим обучить сотрудников по ПК пожарной безопасности для руководителей организаций'])
        self.assertEqual(len(p),1)
    def test_registry_added_numbers(self):
        self.assertTrue(any(x['number']=='1586(1)' for x in cat.catalog()['occupations']))

    def test_price_wins(self):
        for row in (5,6):
            p=cat.result(cat.select_source(f'Экологическая безопасность!B{row}')[0])
            r=schedule([p],'24.08.2026','start')[0]
            self.assertEqual((p['hours'],r[0]['days']),(72,9))
    def test_dpp_not_ot(self):
        import training_rules
        p=cat.result(cat.select_source('Охрана труда!B14')[0])
        self.assertEqual(p['hours'],256)
        self.assertEqual(schedule([p],'24.08.2026','start')[0][0]['days'],32)
        self.assertEqual(training_rules.classify_program(p['name']),training_rules.CAT_PK_DPP)
    def test_round_up(self):
        p=next(cat.result(r) for r in cat.catalog()['programs'] if r['status']=='offered' and not r['hours'])
        for h,days in [(36,5),(178,23),(250,32),(350,44),(520,65),(180,23)]:
            self.assertEqual(schedule([p],'24.08.2026','start',{p['id']:h})[0][0]['days'],days)
    def test_unknowns_fail_closed(self):
        for programs,dates,role in [(cat.resolve_many(None,'Монтажник'),'24.08.2026','start'),(cat.resolve_many('Высота 1 группа'),'24.08.2026','auto'),(cat.resolve_many('Высота 1 группа'),'31.02.2026','start'),(cat.resolve_many('Высота 1 группа'),'01.02.2028','start')]:
            r=schedule(programs,dates,role)
            self.assertEqual(r[0][1],'');self.assertTrue(r[0][0]['schedule_warning'])
    def test_no_compression_or_reversal(self):
        for dates in ['24.08.2026 - 25.08.2026','26.08.2026 - 24.08.2026']:
            self.assertEqual(schedule(cat.resolve_many('Высота 1 группа'),dates)[0][1],'')
    def test_backward(self):
        r=schedule(cat.resolve_many('Высота 2 группа; ОЗП 2 группа'),'31.08.2026','end')
        self.assertEqual(r[0][1],'24.08.2026 - 26.08.2026')
    def test_start_synonyms(self):
        for prefix in ['Дата приема','Дата трудоустройства','Дата выхода на работу','Принят','Приём','Начало обучения']:
            self.assertEqual(parse_dates(prefix+' 24.08.2026')[0],dt.date(2026,8,24))
    def test_header_scope_and_columns(self):
        h=['ФИО','Должность','Дата приема','Дата рождения']
        rows=[h,['Тестов Иван Иванович','Монтажник','24.08.2026','01.01.1990']]
        progs=doc_reader.extract_letter_programs(['Просим обучить сотрудников нашей организации по высоте 2 группа и ОЗП 2 группа'])
        r=doc_reader.extract_students_from_tables([rows],progs)
        self.assertEqual(len(cat.resolve_many(r[0]['program'])),2)
        self.assertEqual(r[0]['date_role'],'start')
    def test_job_not_course(self):
        self.assertFalse(cat.resolve_many(None,'Монтажник технологических трубопроводов 5 разряда')[0]['is_canonical'])
        self.assertFalse(cat.resolve_many('повышение квалификации','Монтажник')[0]['is_canonical'])
    def test_occupation_registry_is_not_offering(self):
        self.assertGreater(len(cat.catalog()['occupations']),5000)
        self.assertFalse(cat.resolve('Авиационный механик по криогенным системам')['is_canonical'])
    def test_demo_excluded(self):
        for status in ['excluded','demo']:
            p=next(r for r in cat.catalog()['programs'] if r['status']==status)
            self.assertFalse(cat.resolve(p['name'])['is_canonical'])
    def test_all_catalog_rows(self):
        for p in cat.catalog()['programs']:
            self.assertTrue(p['source']);self.assertTrue(p['name'])
            self.assertTrue(all(isinstance(h,int) and h>0 for h in p['hours']))
    def test_empty_document_not_fake_student(self):
        import psg_agent
        with patch.object(doc_reader,'parse_raw_text_application',return_value={'students':[]}):
            r=psg_agent.process_application(raw_text='Пусто')
        self.assertFalse(r['success'])
    def test_output_formats(self):
        import docx,openpyxl,psg_agent,contextlib
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            doc=docx.Document();doc.add_paragraph('Просим обучить сотрудников нашей организации по высоте 2 группа и ОЗП 2 группа')
            t=doc.add_table(rows=1,cols=4)
            for c,v in zip(t.rows[0].cells,['ФИО','Должность','Дата приема','Дата рождения']):c.text=v
            for c,v in zip(t.add_row().cells,['Тестов Иван Иванович','Монтажник','24.08.2026','01.01.1990']):c.text=v
            doc.save(root/'source.docx')
            with contextlib.redirect_stdout(io.StringIO()):r=psg_agent.process_application(str(root/'source.docx'),output_file=str(root/'result.docx'))
            self.assertEqual(r['total_enrollments'],2)
            self.assertTrue((root/'result.xlsx').exists())
            wb=openpyxl.load_workbook(root/'result.xlsx');values=[str(c.value) for row in wb.active for c in row]
            self.assertTrue(any('27.08.2026 - 31.08.2026' in v for v in values))
            self.assertTrue(any('24.08.2026 - 26.08.2026' in v for v in values))
    def test_xlsx_header(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            p=Path(temp)/'input.xlsx';w=openpyxl.Workbook();s=w.active
            s.append(['Просим обучить сотрудников по высоте 1 группа'])
            s.append(['ФИО','Дата приема']);s.append(['Тестов Иван Иванович','24.08.2026']);w.save(p)
            r=doc_reader.parse_incoming_application(str(p))
            self.assertEqual(r['students'][0]['date_role'],'start')
            self.assertTrue(cat.resolve_many(r['students'][0]['program'])[0]['is_canonical'])
    def test_text_date_not_birthdate(self):
        r=doc_reader.parse_student_line('Тестов Иван Иванович, Монтажник, дата приема 24.08.2026')
        self.assertEqual(r['birth_date'],'')
        self.assertEqual(r['study_dates'],'Начало обучения 24.08.2026')
    def test_api_catalog_and_validation(self):
        from starlette.testclient import TestClient
        from web_app import app
        client=TestClient(app)
        r=client.get('/api/programs',params={'q':'ОЗП'}).json()
        self.assertTrue(any(p['hours']==[24] for p in r['programs']))
        self.assertTrue(any(p['hours']==[8] for p in r['programs']))
        self.assertEqual(client.post('/api/process',data={'hours_overrides':'[]'}).status_code,422)
        self.assertEqual(client.post('/api/process',data={'date_role':'wrong'}).status_code,422)
    def test_ot_v_not_group_course(self):
        p=cat.result(cat.select_source('Охрана труда!B21')[0])
        self.assertEqual(p['hours'],8)
        import training_rules
        self.assertEqual(training_rules.classify_program(p['name']),training_rules.CAT_OT)

    def test_snils_not_rewritten(self):
        from document_vision import repair_snils_checksum
        s='123-456-789 00';r,ok=repair_snils_checksum(s)
        self.assertEqual(''.join(filter(str.isdigit,r)),''.join(filter(str.isdigit,s)))
        self.assertFalse(ok)

if __name__=='__main__':unittest.main()

"""
Word (.docx) Document Generator for PSG 1C Application Agent.
Creates 1C-compatible Microsoft Word (.docx) applications strictly conforming to the reference structure:
- Page orientation: Landscape (A4, 297 x 210 mm)
- Margins: 0.5 inches (12.7 mm) for maximum readability
- Fonts: Arial 8pt (data) and Arial 8.5pt bold (headers/programs)
- 9 Columns matching reference 1C template:
  1. № п/п
  2. ФИО (в именит. падеже)
  3. ФИО (в дат. падеже)
  4. Должность
  5. Пол
  6. Дата рождения
  7. СНИЛС
  8. Предполагаемые сроки обучения
  9. Электронная почта и телефон обучающегося
- Program divider rows: merged 9 columns, fill #DCFFDD (LIGHT MINT/GREEN как в эталоне 1С)
- Yellow highlighting (#FFFF00) for suspicious / disputed cells matching Excel
- Signature and consent block at bottom
"""

import datetime
from typing import Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from linguistics import clean_application_title

HEADER_TITLES = [
    '№ п/п',
    'ФИО\n(в именит. падеже)',
    'ФИО\n(в дат. падеже)',
    'Должность',
    'Пол',
    'Дата рождения',
    'СНИЛС',
    'Предполагаемые сроки обучения',
    'Электронная почта и телефон обучающегося'
]

# Column widths in inches (total = 10.69 inches, fits A4 Landscape 11.69" with 0.5" margins)
COLUMN_WIDTHS_INCHES = [
    0.55,  # 1. № п/п
    1.75,  # 2. ФИО им.
    1.75,  # 3. ФИО дат.
    1.65,  # 4. Должность
    0.45,  # 5. Пол
    0.95,  # 6. Дата рождения
    1.10,  # 7. СНИЛС
    1.10,  # 8. Сроки
    1.39   # 9. Контакты
]

HEX_MINT_PROGRAM = "DCFFDD"
HEX_YELLOW_DISPUTED = "FFFF00"
HEX_GRAY_HEADER = "F2F4F7"

def _set_cell_background(cell, hex_color: str):
    """Applies background shading color to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for existing_shd in tcPr.findall(docx.oxml.ns.qn('w:shd')):
        tcPr.remove(existing_shd)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def _set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets inner padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for existing_mar in tcPr.findall(docx.oxml.ns.qn('w:tcMar')):
        tcPr.remove(existing_mar)
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def _set_table_borders(table):
    """Applies neat single 0.5pt black/gray borders to the entire table."""
    tblPr = table._tbl.tblPr
    for existing in tblPr.findall(docx.oxml.ns.qn('w:tblBorders')):
        tblPr.remove(existing)
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def _make_row_cant_split(row):
    """Ensures table row does not break awkwardly across page breaks."""
    trPr = row._tr.get_or_add_trPr()
    if not trPr.findall(docx.oxml.ns.qn('w:cantSplit')):
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def _set_repeat_header(row):
    """Marks row as repeating header on every page."""
    trPr = row._tr.get_or_add_trPr()
    if not trPr.findall(docx.oxml.ns.qn('w:tblHeader')):
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def create_1c_application_docx(
    application_title: Optional[str],
    grouped_programs: Dict[str, Dict[str, Any]]
) -> docx.Document:
    """
    Builds the 1C application Microsoft Word (.docx) document.
    - Page setup: A4 Landscape with 0.5" margins
    - Program divider rows have solid LIGHT MINT fill (#DCFFDD).
    - Disputed or suspicious cells are highlighted with solid YELLOW fill (#FFFF00).
    - Signatures and consent block at bottom.
    """
    doc = docx.Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    # 1. Document Title
    if application_title:
        application_title, _ = clean_application_title(application_title)
    else:
        today_str = datetime.date.today().strftime("%d.%m.%Y")
        application_title = f"ЗАЯВКА НА ОБУЧЕНИЕ от {today_str} г."

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run(application_title)
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(11)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0, 0, 0)

    # 2. Add Table
    table = doc.add_table(rows=1, cols=9)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)

    # Header Row
    hdr_row = table.rows[0]
    _make_row_cant_split(hdr_row)
    _set_repeat_header(hdr_row)

    for col_idx, (title_text, width_in) in enumerate(zip(HEADER_TITLES, COLUMN_WIDTHS_INCHES)):
        cell = hdr_row.cells[col_idx]
        cell.width = Inches(width_in)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_background(cell, HEX_GRAY_HEADER)
        _set_cell_margins(cell, top=120, bottom=120, left=80, right=80)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        run = p.add_run(title_text)
        run.font.name = 'Arial'
        run.font.size = Pt(8.5)
        run.bold = True
        run.font.color.rgb = RGBColor(20, 30, 40)

    # 3. Populate Program Groups and Students
    for prog_name, prog_info in grouped_programs.items():
        students = prog_info.get('students', [])
        if not students:
            continue

        is_canonical = prog_info.get('is_canonical', True)
        prog_warning = prog_info.get('warning')
        prog_fill = HEX_YELLOW_DISPUTED if (not is_canonical or prog_warning) else HEX_MINT_PROGRAM

        # Program Divider Row
        prog_row = table.add_row()
        _make_row_cant_split(prog_row)
        
        # Merge all 9 columns
        prog_cell = prog_row.cells[0]
        for c_idx in range(1, 9):
            prog_cell.merge(prog_row.cells[c_idx])
            
        prog_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_background(prog_cell, prog_fill)
        _set_cell_margins(prog_cell, top=100, bottom=100, left=120, right=120)
        
        p_prog = prog_cell.paragraphs[0]
        p_prog.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_prog.paragraph_format.space_before = Pt(0)
        p_prog.paragraph_format.space_after = Pt(0)
        
        run_prog = p_prog.add_run(prog_name)
        run_prog.font.name = 'Arial'
        run_prog.font.size = Pt(8.5)
        run_prog.bold = True
        run_prog.font.color.rgb = RGBColor(0, 0, 0)

        # Student Rows
        for s_idx, student in enumerate(students, start=1):
            s_row = table.add_row()
            _make_row_cant_split(s_row)
            flags = student.get('yellow_flags', {})

            # Prepare values for 9 columns
            row_data = [
                (str(s_idx), WD_ALIGN_PARAGRAPH.RIGHT, WD_CELL_VERTICAL_ALIGNMENT.TOP, False),
                (student.get('fio_nom', ''), WD_ALIGN_PARAGRAPH.LEFT, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'nom_fio' in flags),
                (student.get('fio_dat', ''), WD_ALIGN_PARAGRAPH.LEFT, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'dat_fio' in flags),
                (student.get('position', ''), WD_ALIGN_PARAGRAPH.LEFT, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'position' in flags),
                (student.get('gender', ''), WD_ALIGN_PARAGRAPH.CENTER, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'gender' in flags),
                (student.get('birth_date', ''), WD_ALIGN_PARAGRAPH.CENTER, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'birth_date' in flags),
                (student.get('snils', ''), WD_ALIGN_PARAGRAPH.CENTER, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'snils' in flags),
                (student.get('study_dates', ''), WD_ALIGN_PARAGRAPH.CENTER, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'study_dates' in flags),
                (student.get('contacts', ''), WD_ALIGN_PARAGRAPH.LEFT, WD_CELL_VERTICAL_ALIGNMENT.TOP, 'contacts' in flags)
            ]

            for col_idx, (text_val, align, v_align, is_disputed) in enumerate(row_data):
                cell = s_row.cells[col_idx]
                cell.width = Inches(COLUMN_WIDTHS_INCHES[col_idx])
                cell.vertical_alignment = v_align
                _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
                
                if is_disputed:
                    _set_cell_background(cell, HEX_YELLOW_DISPUTED)

                p = cell.paragraphs[0]
                p.alignment = align
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05

                run = p.add_run(text_val or "")
                run.font.name = 'Arial'
                run.font.size = Pt(8.0)
                run.font.color.rgb = RGBColor(0, 0, 0)

    # 4. Signatures & Consent Block after Table
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    sig_p = doc.add_paragraph()
    sig_p.paragraph_format.space_before = Pt(8)
    sig_p.paragraph_format.space_after = Pt(4)
    sig_p.paragraph_format.line_spacing = 1.15
    
    r_sig1 = sig_p.add_run("Руководитель организации: ____________________ / ____________________ /    М.П.               Дата: «____» ____________ 202_ г.\n")
    r_sig1.font.name = 'Arial'
    r_sig1.font.size = Pt(9.0)
    r_sig1.bold = True

    consent_p = doc.add_paragraph()
    consent_p.paragraph_format.space_before = Pt(4)
    consent_p.paragraph_format.space_after = Pt(0)
    r_consent = consent_p.add_run(
        "Согласие обучающихся на обработку и передачу персональных данных (в т.ч. в ФИС ФРДО и Минтруд РФ) "
        "получено заказчиком в соответствии с Федеральным законом от 27.07.2006 № 152-ФЗ «О персональных данных»."
    )
    r_consent.font.name = 'Arial'
    r_consent.font.size = Pt(7.5)
    r_consent.italic = True
    r_consent.font.color.rgb = RGBColor(100, 100, 100)

    return doc

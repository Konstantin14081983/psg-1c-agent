"""Regression tests use fictional data and mocked external AI only."""
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
import scan_recognition as scan
import ai_vision
import doc_reader
import psg_agent

class ScanRecognition(unittest.TestCase):
    def baseline(self):
        return {'title':'Заявка', 'engine':'OCR', 'students':[
            {'fio_nom':'Тестов Иван Иванович', 'fio_dat':'Тестову Ивану Ивановичу', 'snils':'11223344595', 'program':'ПП', 'preserve_identity':True},
            {'fio_nom':'Примеров Петр Петрович', 'snils':'12345678964', 'program':'ПП', 'preserve_identity':True}]}

    def test_omitted_and_changed_identity_preserves_baseline(self):
        baseline = self.baseline()
        before = copy.deepcopy(baseline)
        result = scan.compare_vision(baseline, {'success':True, 'students':[
            {'fio':'Тестов Иван Иванович', 'snils':'11223344596', 'program':''}]})
        self.assertEqual(result['students'], before['students'])
        self.assertEqual(baseline, before)
        self.assertEqual(len(result['recognition_issues']), 2)
        self.assertIn('snils', result['recognition_issues'][0]['fields'])

    def test_unmatched_ai_rows_remain_review_candidates(self):
        result = scan.compare_vision(self.baseline(), {'success':True, 'students':[{'fio':'Неизвестный Иван Иванович','snils':'999'}]})
        self.assertEqual(len(result['students']),2)
        self.assertEqual(len(result['ai_candidates']),1)
        self.assertTrue(result['recognition_issues'])

    def test_duplicate_rows_are_not_collapsed(self):
        baseline=self.baseline()
        baseline['students'] *= 2
        result=scan.compare_vision(baseline, {'success':True,'students':[]})
        self.assertEqual(len(result['students']),4)

    def test_api_failure_keeps_complete_local_result(self):
        baseline=self.baseline()
        result=scan.compare_vision(baseline, {'success':False,'error':'HTTP 500'})
        self.assertEqual(result['students'],baseline['students'])
        self.assertEqual(result['ai_error'],'HTTP 500')

    def test_normalizer_accepts_program_and_fio_nom(self):
        result=ai_vision._normalize_vision_response({'students':[{'fio_nom':'Тестов Иван Иванович','program':'ПП'}]},'mock')
        self.assertEqual(result['students'][0]['program'],'ПП')
        self.assertEqual(result['students'][0]['fio'],'Тестов Иван Иванович')

    def test_orientation_uses_each_page_osd(self):
        image=Image.new('RGB',(40,80),'white')
        with patch.object(scan,'detect_rotation',return_value=90), patch('document_vision.run_tesseract_on_pil',return_value='Тестов Иван Иванович оказание первой помощи'):
            normalized,text,angle=scan.orient_page(image)
        self.assertEqual(normalized.size,(80,40))
        self.assertEqual(angle,90)

    def test_orientation_fallback_includes_zero_score(self):
        image=Image.new('RGB',(40,80),'white')
        with patch.object(scan,'detect_rotation',return_value=None), patch('document_vision.run_tesseract_on_pil',side_effect=['Иван Иванович помощь','abc','','']):
            _,_,angle=scan.orient_page(image)
        self.assertEqual(angle,0)

    def test_multipage_pdf_same_inputs_and_cleanup(self):
        import fitz
        with tempfile.TemporaryDirectory() as temp:
            pdf=Path(temp)/'scan.pdf'
            with fitz.open() as document:
                document.new_page();document.new_page();document.save(pdf)
            def verify(paths,**kwargs):
                self.assertEqual(len(paths),2)
                self.assertTrue(all(Path(p).exists() for p in paths))
                self.paths=paths[:]
                return {'success':True,'students':[]}
            with patch.object(scan,'orient_page',side_effect=lambda image:(image,'строка',0)), patch('doc_reader.parse_raw_text_application',side_effect=lambda text:copy.deepcopy(self.baseline())), patch('ai_vision.analyze_multi_page_document_with_ai',side_effect=verify):
                local=scan.parse_scan(str(pdf),False)
                ai=scan.parse_scan(str(pdf),True)
            self.assertEqual(local['students'],ai['students'])
            self.assertTrue(all(not Path(p).exists() for p in self.paths))

    def test_manual_start_survives_ai_disagreement_and_exports(self):
        baseline=self.baseline()
        ai=scan.compare_vision(baseline,{'success':True,'students':[]})
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'input.pdf';source.touch()
            results=[]
            for index, parsed in enumerate((baseline,ai)):
                with patch('doc_reader.parse_incoming_application',return_value=copy.deepcopy(parsed)):
                    results.append(psg_agent.process_application(input_file=str(source),output_file=str(Path(temp)/f'out{index}.docx'),manual_overrides={'study_dates':'24.09.2026','date_role':'start'},use_ai=bool(index)))
            def values(result):
                return [{k:v for k,v in row.items() if k!='yellow_flags'} for group in result['grouped_data'].values() for row in group['students']]
            from docx import Document
            import openpyxl
            word_values=[]
            excel_values=[]
            for index in range(2):
                document=Document(str(Path(temp)/f'out{index}.docx'))
                word_values.append([[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables])
                workbook=openpyxl.load_workbook(str(Path(temp)/f'out{index}.xlsx'))
                excel_values.append([list(sheet.values) for sheet in workbook])
                workbook.close()
            self.assertEqual(word_values[0],word_values[1])
            self.assertEqual(excel_values[0],excel_values[1])
            self.assertEqual(values(results[0]),values(results[1]))
            self.assertTrue(all(row['study_dates']=='24.09.2026 - 25.09.2026' for row in values(results[1])))
            self.assertEqual(results[1]['unique_students'],2)
            self.assertTrue(results[1]['requires_review'])
            self.assertTrue(results[1]['audit']['recognition_issues'])

    def test_empty_program_is_not_invented(self):
        text='ФИО дата рождения СНИЛС\nТестов Иван Иванович 01.01.1980 112-233-445 95\nПримеров Петр Петрович 02.02.1981 123-456-789 64'
        rows=doc_reader.parse_scanned_table_text(text)
        self.assertTrue(rows)
        self.assertTrue(all(not row['program'] for row in rows))

    def test_truncated_ai_response_is_not_accepted(self):
        response=type('Response',(),{'status_code':200,'json':lambda self:{'choices':[{'finish_reason':'length','message':{'content':'{"students": []}'}}]}})()
        with patch('ai_vision.get_api_key',return_value='test'), patch('ai_vision.get_base_url',return_value='https://example.invalid/v1'), patch('ai_vision.encode_image_to_base64',return_value=('data:image/png;base64,','image/png')), patch('ai_vision._http_post',return_value=response):
            result=ai_vision.analyze_multi_page_document_with_ai(['a.png','b.png'])
        self.assertFalse(result['success'])
        self.assertIn('обрезан',result['error'])

    def test_mixed_pdf_scanned_continuation_is_not_skipped(self):
        import fitz
        with tempfile.TemporaryDirectory() as temp:
            pdf=Path(temp)/'mixed.pdf'
            image=Path(temp)/'scan.png'
            Image.new('RGB',(50,50),'white').save(image)
            with fitz.open() as document:
                document.new_page().insert_text((50,50),'Vector page')
                page=document.new_page();page.insert_image(page.rect,filename=str(image))
                document.save(pdf)
            with patch.object(scan,'parse_scan',return_value=self.baseline()) as shared:
                result=doc_reader.parse_incoming_application(str(pdf),use_ai=True)
            shared.assert_called_once_with(str(pdf),use_ai=True,api_key=None)
            self.assertEqual(len(result['students']),2)

    def test_tiff_all_frames_and_single_ai_verification(self):
        with tempfile.TemporaryDirectory() as temp:
            image=Path(temp)/'scan.tiff'
            Image.new('RGB',(40,80),'white').save(image,save_all=True,append_images=[Image.new('RGB',(80,40),'white')])
            with patch.object(scan,'orient_page',side_effect=lambda image:(image,'строка',0)) as orient, patch('doc_reader.parse_raw_text_application',side_effect=lambda text:self.baseline()), patch('ai_vision.analyze_multi_page_document_with_ai',return_value={'success':True,'students':[]}) as vision:
                result=doc_reader.parse_incoming_application(str(image),use_ai=True)
            self.assertEqual(orient.call_count,2)
            self.assertEqual(vision.call_count,1)
            self.assertEqual(len(result['students']),2)

    def test_addressee_after_invitation_does_not_replace_first_student(self):
        text='''Просим обучить сотрудников
от Директора
Подписант Николай Иванович
ФИО Дата рождения СНИЛС Должность
Тестов Иван Иванович 01.01.1980 112-233-445 95 Оказание первой помощи
Примеров Петр Петрович 02.02.1981 123-456-789 64 Оказание первой помощи'''
        rows=doc_reader.parse_scanned_table_text(text)
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[0]['fio_nom'],'Тестов Иван Иванович')
